from __future__ import annotations

import csv
import logging
from pathlib import Path
from typing import Optional

from app_config import TABLE_COLUMNS
from models.schemas import PresetEntry

logger = logging.getLogger(__name__)

COLUMN_HEADERS = list(TABLE_COLUMNS)


def _entry_value(entry: PresetEntry, col_name: str) -> str:
    if col_name == "Character Name":
        return entry.character_name
    if col_name == "Character ID":
        return entry.character_id
    if col_name == "Costume Name":
        return entry.costume_name
    if col_name == "Costume Index":
        return str(entry.costume_index)
    if col_name == "Model Preset":
        return str(entry.model_preset)
    if col_name.startswith("Super Skill "):
        try:
            idx = int(col_name.split(" ")[-1]) - 1
        except ValueError:
            return ""
        skills = list(entry.super_skills)
        return skills[idx] if 0 <= idx < len(skills) else ""
    if col_name.startswith("Ultimate Skill "):
        try:
            idx = int(col_name.split(" ")[-1]) - 1
        except ValueError:
            return ""
        ults = list(entry.ultimate_skills)
        return ults[idx] if 0 <= idx < len(ults) else ""
    if col_name == "Awoken Skill":
        return entry.awoken_skill
    if col_name == "Evasive Skill":
        return entry.evasive_skill
    if col_name == "Super Soul":
        return entry.super_soul
    if col_name == "Source":
        return entry.source
    return ""


def _row_for_columns(entry: PresetEntry, columns: list[str]) -> list[str]:
    return [_entry_value(entry, c) for c in columns]


def _truncate_sheet_name(name: str) -> str:
    return name[:31] if name else "Sheet"


def _resolve_columns(columns: Optional[list[str]]) -> list[str]:
    if not columns:
        return list(COLUMN_HEADERS)
    return [c for c in columns if c in COLUMN_HEADERS]


def _preset_header(entry: PresetEntry, columns: list[str]) -> str:
    parts: list[str] = []
    if "Character Name" in columns and entry.character_name:
        parts.append(entry.character_name)
    if "Costume Name" in columns and entry.costume_name:
        parts.append(entry.costume_name)
    if not parts:
        return "Preset"
    if len(parts) == 1:
        return parts[0]
    return f"{parts[0]} — {parts[1]}"


def _md_escape(value: str) -> str:
    return value.replace("|", "\\|").replace("\n", " ").strip()


# ── XLSX ──────────────────────────────────────────────────────────────


def write_xlsx_single(
    path: Path,
    sheets: dict[str, list[PresetEntry]],
    columns: Optional[list[str]] = None,
) -> None:
    import openpyxl

    cols = _resolve_columns(columns)
    wb = openpyxl.Workbook()
    default_sheet = wb.active

    for idx, (sheet_name, entries) in enumerate(sheets.items()):
        ws = default_sheet if idx == 0 else wb.create_sheet()
        ws.title = _truncate_sheet_name(sheet_name)

        for col_idx, col_name in enumerate(cols, start=1):
            ws.cell(row=1, column=col_idx, value=col_name)

        for row_idx, entry in enumerate(entries, start=2):
            for col_idx, value in enumerate(_row_for_columns(entry, cols), start=1):
                ws.cell(row=row_idx, column=col_idx, value=value)

    wb.save(str(path))


def write_xlsx_multi(
    folder: Path,
    sheets: dict[str, list[PresetEntry]],
    columns: Optional[list[str]] = None,
) -> int:
    import openpyxl

    cols = _resolve_columns(columns)
    count = 0
    for sheet_name, entries in sheets.items():
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = _truncate_sheet_name(sheet_name)

        for col_idx, col_name in enumerate(cols, start=1):
            ws.cell(row=1, column=col_idx, value=col_name)

        for row_idx, entry in enumerate(entries, start=2):
            for col_idx, value in enumerate(_row_for_columns(entry, cols), start=1):
                ws.cell(row=row_idx, column=col_idx, value=value)

        wb.save(str(folder / f"{sheet_name}.xlsx"))
        count += 1
    return count


# ── CSV ───────────────────────────────────────────────────────────────


def write_csv_single(
    path: Path,
    sheets: dict[str, list[PresetEntry]],
    columns: Optional[list[str]] = None,
) -> None:
    cols = _resolve_columns(columns)
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        for i, (sheet_name, entries) in enumerate(sheets.items()):
            if i > 0:
                writer.writerow([])
            writer.writerow([f"# Sheet: {sheet_name}"])
            writer.writerow(cols)
            for entry in entries:
                writer.writerow(_row_for_columns(entry, cols))


def write_csv_multi(
    folder: Path,
    sheets: dict[str, list[PresetEntry]],
    columns: Optional[list[str]] = None,
) -> int:
    cols = _resolve_columns(columns)
    count = 0
    for sheet_name, entries in sheets.items():
        path = folder / f"{sheet_name}.csv"
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)
            writer.writerow(cols)
            for entry in entries:
                writer.writerow(_row_for_columns(entry, cols))
        count += 1
    return count


# ── TEXT (Discord / Forum) ────────────────────────────────────────────


def _render_text_sheet(sheet_name: str, entries: list[PresetEntry], cols: list[str], is_multi: bool) -> str:
    lines: list[str] = []
    if is_multi:
        lines.append(f"=== {sheet_name} ===")
        lines.append("")

    for i, entry in enumerate(entries):
        if i > 0:
            lines.append("")
        for col in cols:
            value = _entry_value(entry, col).strip()
            if value:
                lines.append(f"  {col}: {value}")

    return "\n".join(lines)


def write_text_single(
    path: Path,
    sheets: dict[str, list[PresetEntry]],
    columns: Optional[list[str]] = None,
) -> None:
    cols = _resolve_columns(columns)
    multi = len(sheets) > 1
    sections = []
    for sheet_name, entries in sheets.items():
        sections.append(_render_text_sheet(sheet_name, entries, cols, multi))
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n\n".join(sections))


def write_text_multi(
    folder: Path,
    sheets: dict[str, list[PresetEntry]],
    columns: Optional[list[str]] = None,
) -> int:
    cols = _resolve_columns(columns)
    count = 0
    for sheet_name, entries in sheets.items():
        content = _render_text_sheet(sheet_name, entries, cols, is_multi=False)
        with open(folder / f"{sheet_name}.txt", "w", encoding="utf-8") as f:
            f.write(content)
        count += 1
    return count


# ── MARKDOWN ──────────────────────────────────────────────────────────


def _render_md_sheet(sheet_name: str, entries: list[PresetEntry], cols: list[str], is_multi: bool) -> str:
    lines: list[str] = []
    if is_multi:
        lines.append(f"# {sheet_name}")
        lines.append("")

    for i, entry in enumerate(entries):
        if i > 0:
            lines.append("")
        lines.append(f"## {_preset_header(entry, cols)}")
        lines.append("")
        lines.append("| Field | Value |")
        lines.append("|---|---|")
        for col in cols:
            value = _entry_value(entry, col).strip()
            if value:
                lines.append(f"| {_md_escape(col)} | {_md_escape(value)} |")

    return "\n".join(lines)


def write_md_single(
    path: Path,
    sheets: dict[str, list[PresetEntry]],
    columns: Optional[list[str]] = None,
) -> None:
    cols = _resolve_columns(columns)
    multi = len(sheets) > 1
    sections = []
    for sheet_name, entries in sheets.items():
        sections.append(_render_md_sheet(sheet_name, entries, cols, multi))
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n\n".join(sections))


def write_md_multi(
    folder: Path,
    sheets: dict[str, list[PresetEntry]],
    columns: Optional[list[str]] = None,
) -> int:
    cols = _resolve_columns(columns)
    count = 0
    for sheet_name, entries in sheets.items():
        content = _render_md_sheet(sheet_name, entries, cols, is_multi=False)
        with open(folder / f"{sheet_name}.md", "w", encoding="utf-8") as f:
            f.write(content)
        count += 1
    return count